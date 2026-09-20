"""Build the editable Sprint 1 solution overview PowerPoint deck.

Run from the repository root:
    python scripts/create_solution_presentation.py
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


OUT = Path(__file__).resolve().parents[1] / "Docs" / "Simhastha_2028_Sprint1_Solution_Overview.pptx"

# Colour palette adapted from the product design system.
NAVY = "102A43"
TEAL = "007C83"
SAFFRON = "E47B22"
CREAM = "FFF9F1"
INK = "243B53"
MUTED = "627D98"
WHITE = "FFFFFF"
MINT = "D9F3EE"
PALE_ORANGE = "FCE4C5"
PALE_BLUE = "E8F1FA"
LINE = "D9E2EC"
RED = "B42318"


def rgb(hex_code: str) -> RGBColor:
    return RGBColor.from_string(hex_code)


def add_text(slide, text, x, y, w, h, size=18, color=INK, bold=False,
             font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    run = paragraph.runs[0]
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    return box


def shape(slide, kind, x, y, w, h, fill, line=None, radius=True):
    item = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else kind,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    item.fill.solid()
    item.fill.fore_color.rgb = rgb(fill)
    item.line.color.rgb = rgb(line or fill)
    return item


def rect(slide, x, y, w, h, fill, line=None):
    return shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, x, y, w, h, fill, line, False)


def title(slide, heading, eyebrow=None):
    if eyebrow:
        add_text(slide, eyebrow.upper(), 0.65, 0.40, 8.0, 0.25, 10, TEAL, True)
    add_text(slide, heading, 0.65, 0.72, 11.6, 0.58, 28, NAVY, True)
    rect(slide, 0.65, 1.43, 1.02, 0.055, SAFFRON)


def footer(slide, number):
    rect(slide, 0, 7.20, 13.333, 0.30, NAVY)
    add_text(slide, "SIMHASTHA 2028  |  SPRINT 1 FOUNDATION", 0.65, 7.27, 5.3, 0.14, 7.5, WHITE, True)
    add_text(slide, f"{number:02d}", 12.2, 7.25, 0.45, 0.15, 8, WHITE, True, align=PP_ALIGN.RIGHT)


def bullet_list(slide, items, x, y, w, size=15, spacing=0.55, colour=INK):
    for index, item in enumerate(items):
        add_text(slide, "•", x, y + index * spacing, 0.22, 0.25, size + 1, SAFFRON, True)
        add_text(slide, item, x + 0.27, y + index * spacing, w - 0.27, 0.40, size, colour)


def card(slide, heading, body, x, y, w, h, accent=TEAL):
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h, WHITE, LINE)
    rect(slide, x, y, 0.08, h, accent)
    add_text(slide, heading, x + 0.25, y + 0.22, w - 0.45, 0.28, 15, NAVY, True)
    add_text(slide, body, x + 0.25, y + 0.62, w - 0.45, h - 0.80, 11.5, MUTED)


def icon_tile(slide, symbol, heading, body, x, y, fill):
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, 2.66, 1.55, fill, fill)
    add_text(slide, symbol, x + 0.18, y + 0.22, 0.42, 0.35, 21, TEAL if fill != PALE_ORANGE else SAFFRON, True)
    add_text(slide, heading, x + 0.68, y + 0.23, 1.75, 0.25, 13, NAVY, True)
    add_text(slide, body, x + 0.20, y + 0.73, 2.22, 0.55, 10.5, INK)


def add_slide_numbered(prs, number):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0, 0, 13.333, 7.5, CREAM)
    footer(slide, number)
    return slide


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "Simhastha 2028 - Sprint 1 Solution Overview"
    prs.core_properties.subject = "Digital Experience Platform prototype"
    prs.core_properties.author = "Simhastha 2028 project team"

    # 1. Cover
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0, 0, 13.333, 7.5, NAVY)
    rect(slide, 0, 0, 0.18, 7.5, SAFFRON)
    shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, 9.5, 0.5, 3.0, 3.0, TEAL, TEAL)
    shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, 10.2, 3.0, 2.55, 2.55, SAFFRON, SAFFRON)
    shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, 8.85, 4.55, 1.45, 1.45, MINT, MINT)
    add_text(slide, "SIMHASTHA 2028", 0.75, 0.85, 4.5, 0.25, 13, "8DE1DD", True)
    add_text(slide, "A safer digital\nexperience for Ujjain", 0.75, 1.40, 7.8, 1.45, 32, WHITE, True)
    add_text(slide, "Sprint 1 solution overview | Public information, maps, CMS and simulated support flows", 0.77, 3.15, 6.8, 0.65, 16, "D9E2EC")
    rect(slide, 0.75, 4.20, 1.30, 0.07, SAFFRON)
    add_text(slide, "Foundation prototype • Local Docker deployment • EN + HI scaffold", 0.77, 4.48, 6.6, 0.28, 12, WHITE, True)
    add_text(slide, "Solution brief", 0.77, 6.70, 2.0, 0.22, 10, "8DE1DD", True)

    # 2. Opportunity
    slide = add_slide_numbered(prs, 2)
    title(slide, "The challenge: high-stakes information, fragmented journeys", "Why this matters")
    card(slide, "For pilgrims", "Planning a visit needs one clear place for events, sacred sites, ghats and essential guidance — across languages and devices.", 0.75, 1.85, 3.85, 2.0, TEAL)
    card(slide, "For administrators", "Timely public content needs a controlled workflow, clear roles, and a way to publish updates without code changes.", 4.75, 1.85, 3.85, 2.0, SAFFRON)
    card(slide, "For public safety", "Digital features must be useful without implying that a prototype is connected to real dispatch, CCTV, or government systems.", 8.75, 1.85, 3.85, 2.0, RED)
    add_text(slide, "Design principle", 0.75, 4.45, 1.8, 0.25, 12, TEAL, True)
    add_text(slide, "Make verified information easy to find — and make prototype limits impossible to miss.", 0.75, 4.80, 11.0, 0.50, 23, NAVY, True)
    add_text(slide, "The result is an MVP that improves discoverability today while creating a safe base for operational capabilities later.", 0.75, 5.60, 10.8, 0.35, 14, MUTED)

    # 3. Solution
    slide = add_slide_numbered(prs, 3)
    title(slide, "One platform, designed around the pilgrim journey", "The solution")
    add_text(slide, "Discover", 1.02, 1.78, 1.5, 0.25, 15, NAVY, True, align=PP_ALIGN.CENTER)
    add_text(slide, "Navigate", 3.72, 1.78, 1.5, 0.25, 15, NAVY, True, align=PP_ALIGN.CENTER)
    add_text(slide, "Get support", 6.42, 1.78, 1.5, 0.25, 15, NAVY, True, align=PP_ALIGN.CENTER)
    add_text(slide, "Stay informed", 9.12, 1.78, 1.7, 0.25, 15, NAVY, True, align=PP_ALIGN.CENTER)
    for x, fill, symbol in [(1.18, PALE_ORANGE, "01"), (3.88, PALE_BLUE, "02"), (6.58, MINT, "03"), (9.28, PALE_ORANGE, "04")]:
        shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x, 2.18, 1.15, 1.15, fill, fill)
        add_text(slide, symbol, x, 2.54, 1.15, 0.25, 16, NAVY, True, align=PP_ALIGN.CENTER)
    for x in [2.5, 5.2, 7.9]:
        rect(slide, x, 2.70, 0.55, 0.05, TEAL)
    add_text(slide, "Events, temples\nand ghats", 0.80, 3.62, 1.95, 0.55, 13, INK, align=PP_ALIGN.CENTER)
    add_text(slide, "Interactive map\nand place layers", 3.50, 3.62, 1.95, 0.55, 13, INK, align=PP_ALIGN.CENTER)
    add_text(slide, "Directory, SOS,\ncase reporting", 6.20, 3.62, 1.95, 0.55, 13, INK, align=PP_ALIGN.CENTER)
    add_text(slide, "Announcements\nand CMS updates", 8.90, 3.62, 1.95, 0.55, 13, INK, align=PP_ALIGN.CENTER)
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 1.18, 4.95, 10.95, 0.78, NAVY, NAVY)
    add_text(slide, "Shared foundations: mobile-first UI  •  English/Hindi scaffold  •  accessibility-aware components  •  role-based CMS", 1.35, 5.23, 10.6, 0.25, 14, WHITE, True, align=PP_ALIGN.CENTER)

    # 4. Experience modules
    slide = add_slide_numbered(prs, 4)
    title(slide, "A practical MVP for public discovery and support", "What users can do")
    icon_tile(slide, "◷", "Calendar & events", "Browse CMS-sourced event listings with category filters and freshness timestamps.", 0.75, 1.75, PALE_BLUE)
    icon_tile(slide, "⌖", "Map & locations", "View temple, ghat and emergency layers on a MapLibre + OpenStreetMap base.", 3.62, 1.75, MINT)
    icon_tile(slide, "⌂", "Sacred-site directory", "Explore nine named temple profiles and basic ghat information.", 6.49, 1.75, PALE_ORANGE)
    icon_tile(slide, "!", "Emergency directory", "Find clearly marked demo directory entries and a simulated SOS flow.", 9.36, 1.75, "FDECEC")
    icon_tile(slide, "↗", "Lost & Found", "Submit a report and check progress using a private, opaque reference.", 2.19, 3.72, MINT)
    icon_tile(slide, "♥", "Missing-person reports", "Submit a consent-based report; public lookup exposes only status and timestamps.", 5.06, 3.72, "FDECEC")
    icon_tile(slide, "⚙", "Admin CMS", "Login, manage content, and move events, temples, ghats and announcements through publish states.", 7.93, 3.72, PALE_BLUE)
    add_text(slide, "Persistent emergency access and mobile-friendly forms are built into the public experience.", 0.75, 6.10, 11.4, 0.28, 14, MUTED, align=PP_ALIGN.CENTER)

    # 5. Safety and privacy
    slide = add_slide_numbered(prs, 5)
    title(slide, "Safety and privacy are product features — not footnotes", "Trust by design")
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 0.75, 1.75, 3.55, 3.95, "FDECEC", "F3B6B1")
    add_text(slide, "SIMULATED,\nNOT DISPATCHED", 1.05, 2.10, 2.95, 0.75, 22, RED, True, align=PP_ALIGN.CENTER)
    add_text(slide, "SOS records stay inside the prototype. No police, ambulance, fire, CCTV or government system is notified.", 1.05, 3.25, 2.95, 1.0, 15, INK, align=PP_ALIGN.CENTER)
    add_text(slide, "Visible safety banners appear on every critical screen and response.", 1.05, 4.80, 2.95, 0.42, 12, RED, True, align=PP_ALIGN.CENTER)
    card(slide, "Data minimisation", "Missing-person public status checks never expose a name, photo, location, contact details or search/listing surface.", 4.75, 1.75, 3.55, 1.70, TEAL)
    card(slide, "Opaque case references", "Cryptographically random references support follow-up without sequential IDs that could enable case enumeration.", 8.75, 1.75, 3.55, 1.70, SAFFRON)
    card(slide, "Honest prototype data", "Seeded emergency numbers are non-routable placeholders. Temple and event information is flagged as placeholder until verified.", 4.75, 4.00, 3.55, 1.70, RED)
    card(slide, "Controlled access", "JWT authentication and role-based admin access separate public viewing from content operations and case verification.", 8.75, 4.00, 3.55, 1.70, TEAL)

    # 6. Architecture
    slide = add_slide_numbered(prs, 6)
    title(slide, "A modular, locally runnable architecture", "How it works")
    layers = [
        ("Public web + Admin CMS", "Next.js • React • TypeScript • Tailwind • next-intl • MapLibre", TEAL),
        ("Application API", "FastAPI • Pydantic validation • JWT/RBAC • OpenAPI", NAVY),
        ("Data & services", "PostgreSQL + PostGIS • Redis • Alembic migrations • audit log", SAFFRON),
    ]
    y = 1.75
    for name, desc, fill in layers:
        shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 1.00, y, 7.35, 0.90, fill, fill)
        add_text(slide, name, 1.30, y + 0.18, 2.3, 0.22, 15, WHITE, True)
        add_text(slide, desc, 3.62, y + 0.20, 4.35, 0.24, 12, WHITE)
        y += 1.12
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 9.05, 1.75, 3.25, 3.15, WHITE, LINE)
    add_text(slide, "LOCAL DEV STACK", 9.35, 2.07, 2.65, 0.22, 13, TEAL, True, align=PP_ALIGN.CENTER)
    add_text(slide, "Docker Compose", 9.35, 2.63, 2.65, 0.25, 20, NAVY, True, align=PP_ALIGN.CENTER)
    add_text(slide, "PostGIS  •  Redis\nFastAPI  •  Next.js", 9.35, 3.20, 2.65, 0.55, 13, MUTED, align=PP_ALIGN.CENTER)
    add_text(slide, "One-command start/stop scripts for local operation", 9.35, 4.12, 2.65, 0.38, 11, INK, True, align=PP_ALIGN.CENTER)
    add_text(slide, "Designed as a foundation: clear API contracts and separable services make later integrations easier to govern.", 1.00, 5.55, 10.9, 0.40, 15, MUTED, align=PP_ALIGN.CENTER)

    # 7. Admin workflow
    slide = add_slide_numbered(prs, 7)
    title(slide, "A content operating model, not hard-coded pages", "CMS and governance")
    add_text(slide, "1", 1.08, 2.12, 0.55, 0.35, 24, WHITE, True, align=PP_ALIGN.CENTER)
    shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, 0.95, 1.90, 0.82, 0.82, TEAL, TEAL)
    add_text(slide, "Create", 0.78, 2.98, 1.15, 0.25, 14, NAVY, True, align=PP_ALIGN.CENTER)
    add_text(slide, "Events, announcements, temples and ghats", 0.52, 3.32, 1.70, 0.65, 11, MUTED, align=PP_ALIGN.CENTER)
    for x in [2.22, 4.72, 7.22]:
        rect(slide, x, 2.27, 1.25, 0.06, LINE)
    for idx, label, desc, x, colour in [
        ("2", "Review", "Role-gated access\nfor content teams", 3.10, SAFFRON),
        ("3", "Publish", "Draft → published\nworkflow", 5.60, TEAL),
        ("4", "Audit", "Sensitive-case and\npublish actions logged", 8.10, NAVY),
    ]:
        shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x, 1.90, 0.82, 0.82, colour, colour)
        add_text(slide, idx, x + 0.13, 2.12, 0.55, 0.35, 24, WHITE, True, align=PP_ALIGN.CENTER)
        add_text(slide, label, x - 0.17, 2.98, 1.16, 0.25, 14, NAVY, True, align=PP_ALIGN.CENTER)
        add_text(slide, desc, x - 0.38, 3.32, 1.60, 0.65, 11, MUTED, align=PP_ALIGN.CENTER)
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 0.90, 4.75, 11.45, 0.84, PALE_BLUE, PALE_BLUE)
    add_text(slide, "Roles available today", 1.25, 5.02, 2.1, 0.20, 13, NAVY, True)
    add_text(slide, "Super Admin", 3.65, 5.02, 1.2, 0.20, 13, TEAL, True)
    add_text(slide, "Content Manager", 5.45, 5.02, 1.65, 0.20, 13, TEAL, True)
    add_text(slide, "Public User", 7.80, 5.02, 1.2, 0.20, 13, TEAL, True)
    add_text(slide, "Role model is structured to extend for operations teams later.", 1.25, 5.34, 9.7, 0.18, 11, MUTED)

    # 8. Evidence
    slide = add_slide_numbered(prs, 8)
    title(slide, "Foundation delivered: testable, working MVP slices", "Sprint 1 evidence")
    stats = [("9", "named temple\nprofiles"), ("75", "backend tests\npassing"), ("4", "local Docker\nservices"), ("2", "locales in the\ni18n scaffold")]
    x = 0.90
    for value, label in stats:
        shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, 1.80, 2.65, 1.45, WHITE, LINE)
        add_text(slide, value, x, 2.04, 2.65, 0.42, 28, TEAL, True, align=PP_ALIGN.CENTER)
        add_text(slide, label, x, 2.56, 2.65, 0.42, 11, MUTED, align=PP_ALIGN.CENTER)
        x += 2.92
    card(slide, "Public API coverage", "Events • temples • ghats • emergency directory • SOS • lost & found • missing-person status lookups", 0.90, 4.00, 3.55, 1.40, TEAL)
    card(slide, "Quality guardrails", "Input validation, access-control checks, privacy regression tests, audit records and safety-notice assertions.", 4.88, 4.00, 3.55, 1.40, SAFFRON)
    card(slide, "Developer experience", "Docker Compose stack, migrations, idempotent seed data, API docs, plus start/stop scripts for local operation.", 8.86, 4.00, 3.55, 1.40, NAVY)

    # 9. Boundaries
    slide = add_slide_numbered(prs, 9)
    title(slide, "What this solution deliberately does not claim", "Responsible scope")
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 0.75, 1.75, 5.55, 3.85, "FDECEC", "F3B6B1")
    add_text(slide, "NOT LIVE OPERATIONS", 1.05, 2.08, 4.95, 0.28, 19, RED, True, align=PP_ALIGN.CENTER)
    bullet_list(slide, ["No live emergency dispatch or responder notification", "No real emergency numbers or government-system connection", "No CCTV, biometrics, facial recognition or crowd sensors", "No fabricated crowd, traffic, parking or verified timetable data"], 1.10, 2.72, 4.80, 13.5, 0.58)
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 6.95, 1.75, 5.55, 3.85, MINT, "A8DED5")
    add_text(slide, "A FOUNDATION FOR NEXT", 7.25, 2.08, 4.95, 0.28, 19, TEAL, True, align=PP_ALIGN.CENTER)
    bullet_list(slide, ["Governed integration with approved departments and data owners", "Verified content and multilingual review processes", "Live-map, routing, transport and parking capabilities", "Operational workflows only after policy, safety and procurement readiness"], 7.30, 2.72, 4.80, 13.5, 0.58)
    add_text(slide, "Clear boundaries protect users now and prevent a prototype from being mistaken for an operational system.", 0.90, 6.05, 11.5, 0.35, 15, NAVY, True, align=PP_ALIGN.CENTER)

    # 10. Roadmap
    slide = add_slide_numbered(prs, 10)
    title(slide, "A staged path from foundation to operations", "Roadmap")
    phases = [
        ("Sprint 1", "Foundation", "Public site • CMS • master data • map • simulated support", TEAL),
        ("Phase 2", "Pilgrim platform", "Trip planner • live-map upgrades • parking • transport • PWA • AI assistant", SAFFRON),
        ("Phase 3", "Operations", "Command centre • crowd analytics • traffic • incident workflows", NAVY),
        ("Phase 4", "Event operations", "Real-time monitoring • alerts • operational analytics", "506D85"),
    ]
    x = 0.75
    for phase, name, desc, colour in phases:
        shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, 1.85, 2.85, 3.45, WHITE, LINE)
        rect(slide, x, 1.85, 2.85, 0.18, colour)
        add_text(slide, phase.upper(), x + 0.22, 2.28, 2.4, 0.20, 10, colour, True)
        add_text(slide, name, x + 0.22, 2.67, 2.4, 0.32, 17, NAVY, True)
        add_text(slide, desc, x + 0.22, 3.35, 2.38, 1.05, 12, MUTED)
        x += 3.10
    add_text(slide, "Progression principle: add real-time capabilities only with authoritative data, explicit governance, and operational ownership.", 0.85, 5.93, 11.55, 0.38, 15, NAVY, True, align=PP_ALIGN.CENTER)

    # 11. Close
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0, 0, 13.333, 7.5, NAVY)
    rect(slide, 0, 0, 13.333, 0.18, SAFFRON)
    add_text(slide, "THE TAKEAWAY", 0.75, 0.95, 3.0, 0.25, 12, "8DE1DD", True)
    add_text(slide, "A credible digital starting point\nfor Simhastha 2028", 0.75, 1.50, 8.1, 1.25, 31, WHITE, True)
    add_text(slide, "The platform makes important information easier to discover, gives content teams a practical operating layer, and treats safety and privacy as non-negotiable design constraints.", 0.77, 3.25, 7.7, 0.78, 16, "D9E2EC")
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 0.77, 4.62, 7.4, 0.72, TEAL, TEAL)
    add_text(slide, "Next: validate content owners, integration governance, and Phase 2 priorities.", 1.00, 4.86, 6.95, 0.25, 15, WHITE, True, align=PP_ALIGN.CENTER)
    shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, 9.35, 1.30, 2.55, 2.55, TEAL, TEAL)
    shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, 10.20, 3.35, 1.75, 1.75, SAFFRON, SAFFRON)
    shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, 8.85, 4.65, 1.15, 1.15, MINT, MINT)
    add_text(slide, "Thank you", 0.77, 6.55, 1.7, 0.24, 12, "8DE1DD", True)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"Created {OUT}")


if __name__ == "__main__":
    build()
